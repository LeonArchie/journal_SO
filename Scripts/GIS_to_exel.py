import pandas as pd
import re
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
import traceback
import sys
from datetime import datetime
import csv

def log_message(message):
    """Запись сообщения в лог-файл"""
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    log_entry = f"[{timestamp}] {message}\n"
    print(message)
    with open('GIS_to_EXEL.txt', 'a', encoding='utf-8') as log_file:
        log_file.write(log_entry)

def parse_schedule_file(filename):
    """Парсинг исходного файла с расписанием с использованием CSV reader"""
    log_message(f"Начинаем парсинг файла: {filename}")
    classes = {}
    current_class = None
    current_schedule = []
    
    try:
        with open(filename, 'r', encoding='windows-1251') as file:
            # Используем CSV reader для правильного парсинга кавычек
            csv_reader = csv.reader(file, delimiter=';', quotechar='"')
            
            for row_num, row in enumerate(csv_reader, 1):
                # Пропускаем пустые строки
                if not row or all(cell.strip() == '' for cell in row):
                    continue
                    
                # Объединяем ячейки в строку для анализа
                line_content = ';'.join(row)
                
                # Определяем начало нового класса
                if line_content.startswith('Класс:'):
                    if current_class and current_schedule:
                        classes[current_class] = current_schedule
                        log_message(f"  Сохранен класс {current_class} с {len(current_schedule)} строками расписания")
                    
                    current_class = line_content.split(':')[1].strip()
                    current_schedule = []
                    log_message(f"Найден новый класс: {current_class}")
                    continue
                    
                # Пропускаем заголовки дней недели
                if line_content.startswith(';Пн;Вт;Ср;Чт;Пт'):
                    log_message(f"  Пропущен заголовок дней недели в строке {row_num}")
                    continue
                
                # Добавляем строку расписания (начинается с ;)
                if len(row) > 1 and row[0] == '':
                    current_schedule.append(row)
                    if len(current_schedule) % 10 == 0:  # Логируем каждые 10 строк
                        log_message(f"  Добавлено {len(current_schedule)} строк для класса {current_class}")
    
    except Exception as e:
        log_message(f"ОШИБКА при парсинге файла на строке {row_num}: {e}")
        log_message(f"Трассировка: {traceback.format_exc()}")
        raise
    
    # Добавляем последний класс
    if current_class and current_schedule:
        classes[current_class] = current_schedule
        log_message(f"  Сохранен последний класс {current_class} с {len(current_schedule)} строками расписания")
    
    log_message(f"Парсинг завершен. Найдено классов: {len(classes)}")
    return classes

def process_schedule_data(classes_data):
    """Обработка данных расписания и создание структуры для Excel"""
    log_message("Начинаем обработку данных расписания")
    all_data = {}
    
    for class_name, schedule_rows in classes_data.items():
        log_message(f"Обрабатываем класс: {class_name}")
        log_message(f"  Количество строк расписания: {len(schedule_rows)}")
        
        # Создаем данные для DataFrame
        data = []
        lesson_numbers = []
        
        for row_num, row in enumerate(schedule_rows, 1):
            log_message(f"  Строка {row_num}: содержит {len(row)} ячеек")
            
            # Обрабатываем строку расписания
            if len(row) >= 2:
                # Первая ячейка - пустая, вторая - номер урока
                lesson_num = row[1] if len(row) > 1 else ''
                lesson_numbers.append(lesson_num)
                
                # Извлекаем данные для каждого дня недели
                # row[2] - Пн, row[3] - Вт, row[4] - Ср, row[5] - Чт, row[6] - Пт
                days_data = []
                for i in range(2, 7):  # Индексы 2-6 для дней недели
                    if i < len(row):
                        days_data.append(row[i])
                    else:
                        days_data.append('')  # Заполняем пустыми строками если не хватает данных
                
                row_data = [lesson_num] + days_data
                log_message(f"    Номер урока: '{lesson_num}', данные дней: {[len(d) if d else 0 for d in days_data]}")
                data.append(row_data)
            else:
                log_message(f"    ПРЕДУПРЕЖДЕНИЕ: строка содержит только {len(row)} ячеек")
        
        log_message(f"  Собрано {len(data)} строк данных для DataFrame")
        
        try:
            # Создаем DataFrame
            df = pd.DataFrame(data, columns=['Урок', 'Пн', 'Вт', 'Ср', 'Чт', 'Пт'])
            log_message(f"  DataFrame создан успешно. Размер: {df.shape}")
            
            # Логируем первые несколько строк для проверки
            for i in range(min(3, len(df))):
                log_message(f"    Строка {i+1}: Урок='{df.iloc[i]['Урок']}', Пн='{df.iloc[i]['Пн'][:50] if df.iloc[i]['Пн'] else ''}...'")
            
            all_data[class_name] = {
                'dataframe': df,
                'lesson_numbers': lesson_numbers
            }
            
        except Exception as e:
            log_message(f"  ОШИБКА при создании DataFrame для класса {class_name}: {e}")
            log_message(f"  Размер данных: {len(data)} строк")
            if data:
                log_message(f"  Первая строка данных: {data[0]}")
                log_message(f"  Количество элементов в первой строке: {len(data[0])}")
            raise
    
    log_message(f"Обработка данных завершена. Обработано классов: {len(all_data)}")
    return all_data

def clean_cell_value(value):
    """Очистка значения ячейки от лишних пробелов и переносов строк"""
    if value is None:
        return ''
    value = str(value).strip()
    # Заменяем множественные переносы строк на один
    value = re.sub(r'\n+', '\n', value)
    return value

def create_excel_with_merged_lessons(classes_data, output_filename):
    """Создание Excel файла с объединенными ячейками уроков"""
    log_message(f"Начинаем создание Excel файла: {output_filename}")
    wb = Workbook()
    
    # Удаляем стандартный лист
    wb.remove(wb.active)
    log_message("Стандартный лист удален")
    
    for class_name, class_info in classes_data.items():
        log_message(f"Создаем лист для класса: {class_name}")
        df = class_info['dataframe']
        
        # Создаем новый лист
        sheet_name = class_name[:31]  # Ограничение длины имени листа в Excel
        ws = wb.create_sheet(title=sheet_name)
        log_message(f"  Лист '{sheet_name}' создан")
        
        # Добавляем заголовок класса
        ws.merge_cells('A1:F1')
        ws['A1'] = f'Класс: {class_name}'
        ws['A1'].alignment = Alignment(horizontal='center')
        log_message("  Заголовок класса добавлен")
        
        # Добавляем заголовки дней недели
        headers = ['Урок', 'Пн', 'Вт', 'Ср', 'Чт', 'Пт']
        for col, header in enumerate(headers, 1):
            ws.cell(row=3, column=col, value=header)
        log_message("  Заголовки дней недели добавлены")
        
        # Добавляем данные
        row_count = 0
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), 4):
            row_count += 1
            for c_idx, value in enumerate(row, 1):
                cleaned_value = clean_cell_value(value)
                ws.cell(row=r_idx, column=c_idx, value=cleaned_value)
        log_message(f"  Данные добавлены: {row_count} строк")
        
        # Объединяем ячейки с номерами уроков и заменяем значения
        log_message("  Начинаем объединение ячеек с номерами уроков")
        
        # Сначала собираем информацию о всех уроках и их позициях
        lesson_groups = {}
        
        for row_idx in range(4, len(df) + 4):
            lesson_cell = ws.cell(row=row_idx, column=1)
            lesson_value = lesson_cell.value
            
            if lesson_value:
                # Определяем основной номер урока
                if re.match(r'^\d+$', str(lesson_value)):
                    main_lesson = lesson_value
                elif re.match(r'^\d+\.\d+$', str(lesson_value)):
                    main_lesson = re.match(r'^(\d+)', str(lesson_value)).group(1)
                else:
                    continue  # Пропускаем невалидные значения
                
                # Добавляем позицию в группу
                if main_lesson not in lesson_groups:
                    lesson_groups[main_lesson] = []
                lesson_groups[main_lesson].append(row_idx)
        
        # Объединяем ячейки для каждой группы и заменяем значения
        merge_count = 0
        for main_lesson, row_indices in lesson_groups.items():
            if len(row_indices) > 1:
                # Сортируем позиции по возрастанию
                row_indices.sort()
                start_row = row_indices[0]
                end_row = row_indices[-1]
                
                # Заменяем ВСЕ значения в группе на основной номер урока
                for row_idx in row_indices:
                    ws.cell(row=row_idx, column=1, value=main_lesson)
                
                # Объединяем ячейки
                ws.merge_cells(f'A{start_row}:A{end_row}')
                merge_count += 1
                log_message(f"    Объединены ячейки A{start_row}:A{end_row} для урока {main_lesson}")
                
                # Центрируем текст в объединенной ячейке
                ws.cell(row=start_row, column=1).alignment = Alignment(
                    horizontal='center', vertical='center'
                )
            else:
                # Для одиночных уроков тоже заменяем значение если это подурок
                row_idx = row_indices[0]
                current_value = ws.cell(row=row_idx, column=1).value
                if re.match(r'^\d+\.\d+$', str(current_value)):
                    ws.cell(row=row_idx, column=1, value=main_lesson)
                    log_message(f"    Заменено значение '{current_value}' на '{main_lesson}' в строке {row_idx}")
        
        log_message(f"  Объединено {merge_count} групп ячеек")
        
        # Настраиваем ширину колонок и перенос текста
        for col_idx in range(1, 7):  # Колонки A-F
            column_letter = get_column_letter(col_idx)
            max_length = 0
            
            # Проходим по всем строкам в колонке, пропуская объединенные ячейки
            for row_idx in range(1, len(df) + 4):
                try:
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if cell.value and hasattr(cell, 'value'):
                        # Для ячеек с данными устанавливаем перенос текста
                        cell.alignment = Alignment(wrap_text=True, vertical='top')
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    # Пропускаем объединенные ячейки
                    continue
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
            log_message(f"    Колонка {column_letter} установлена ширина: {adjusted_width}")
    
    # Сохраняем файл
    try:
        wb.save(output_filename)
        log_message(f"Файл {output_filename} успешно создан!")
        log_message("="*50)
    except Exception as e:
        log_message(f"ОШИБКА при сохранении файла: {e}")
        raise

def main():
    # Очищаем лог-файл при каждом запуске
    with open('GIS_to_EXEL.txt', 'w', encoding='utf-8') as log_file:
        log_file.write(f"Лог преобразования CSV в Excel\n")
        log_file.write(f"Запуск: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        log_file.write("="*50 + "\n")
    
    input_filename = 'GIS_schedule.csv'
    output_filename = 'GIS.xlsx'
    
    try:
        log_message("Начало выполнения скрипта")
        
        # Парсим исходный файл
        classes_data = parse_schedule_file(input_filename)
        log_message(f"Найдено классов: {len(classes_data)}")
        
        # Обрабатываем данные
        processed_data = process_schedule_data(classes_data)
        
        # Создаем Excel файл с объединенными ячейками
        create_excel_with_merged_lessons(processed_data, output_filename)
        
        log_message("Скрипт выполнен успешно!")
        
    except Exception as e:
        log_message(f"КРИТИЧЕСКАЯ ОШИБКА: {e}")
        log_message(f"Трассировка: {traceback.format_exc()}")
        sys.exit(1)

if __name__ == "__main__":
    main()